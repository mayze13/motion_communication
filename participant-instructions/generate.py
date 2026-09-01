"""Generate per-participant maze-instruction handouts (.docx) for Minds in Motion.

Usage:
    pip install -r requirements.txt
    python generate.py

Produces output/participant_001.docx .. participant_250.docx, a combined
output/all_participants.docx (one page per participant, in the same order),
plus sequences.csv (the machine-readable version of each participant's
sequence).
"""
import csv
import random
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

FIRST_ID = 1
LAST_ID = 250
SEQUENCE_LENGTH = 100
NUMBER_MIN, NUMBER_MAX = 1, 12
GRID_COLS = 10  # SEQUENCE_LENGTH must be divisible by GRID_COLS

UCL_PURPLE = RGBColor(0x36, 0x1A, 0x54)
ROW_TINT = "F5F1FA"  # light UCL-purple tint, alternating rows
FONT_NAME = "Arial"  # print-safe; not guaranteed the printing machine has DM Sans

TITLE_TEXT = "Participant Instructions"
INSTRUCTIONS_TEXT = (
    "Please navigate through the maze to each of the following screens in "
    "order. Read the numbers from left to right. You may use the pen ONLY "
    "to tick off the numbers."
)

OUTPUT_DIR = Path(__file__).parent / "output"
SEQUENCES_CSV = Path(__file__).parent / "sequences.csv"


def sequence_for(participant_id: int) -> list[int]:
    """Build the 100-number sequence out of shuffled 1..12 blocks.

    Each block of up to 12 is a random permutation of the waypoints (all 12
    are visited once before any repeats within that block), and the last
    element of one block is never equal to the first element of the next,
    so the same number never appears twice in a row even across the
    block boundary.
    """
    rng = random.Random(participant_id)
    waypoints = list(range(NUMBER_MIN, NUMBER_MAX + 1))

    sequence: list[int] = []
    prev_last = None
    remaining = SEQUENCE_LENGTH
    while remaining > 0:
        block_size = min(len(waypoints), remaining)
        block = rng.sample(waypoints, block_size)
        if prev_last is not None and block[0] == prev_last and block_size > 1:
            swap_idx = next(i for i in range(1, block_size) if block[i] != prev_last)
            block[0], block[swap_idx] = block[swap_idx], block[0]
        sequence.extend(block)
        prev_last = block[-1]
        remaining -= block_size

    return sequence


def set_cell_shading(cell, hex_color: str) -> None:
    shd = cell._tc.get_or_add_tcPr().makeelement(qn("w:shd"), {
        qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): hex_color,
    })
    cell._tc.get_or_add_tcPr().append(shd)


def set_cell_borders(cell) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.makeelement(qn("w:tcBorders"), {})
    for edge in ("top", "left", "bottom", "right"):
        el = borders.makeelement(qn(f"w:{edge}"), {
            qn("w:val"): "single", qn("w:sz"): "6",
            qn("w:space"): "0", qn("w:color"): "8F6FC0",
        })
        borders.append(el)
    tcPr.append(borders)


def style_run(run, size=11, bold=False, color=None):
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def setup_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    doc.styles["Normal"].font.name = FONT_NAME
    doc.styles["Normal"].font.size = Pt(11)


def add_participant_content(doc: Document, participant_id: int, sequence: list[int]) -> None:
    """Append one participant's title/instructions/grid to doc as a single page."""
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(title.add_run(TITLE_TEXT), size=26, bold=True, color=UCL_PURPLE)

    pid = doc.add_paragraph()
    pid.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pid.space_after = Pt(18)
    style_run(pid.add_run(f"Participant ID: {participant_id}"), size=16, bold=True)

    instructions = doc.add_paragraph()
    instructions.space_after = Pt(18)
    style_run(instructions.add_run(INSTRUCTIONS_TEXT), size=12)

    rows = SEQUENCE_LENGTH // GRID_COLS
    table = doc.add_table(rows=rows, cols=GRID_COLS)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_width = Cm(1.55)
    for i, number in enumerate(sequence):
        r, c = divmod(i, GRID_COLS)
        cell = table.cell(r, c)
        cell.width = col_width
        set_cell_borders(cell)
        if r % 2 == 1:
            set_cell_shading(cell, ROW_TINT)

        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(para.add_run(str(number)), size=16, bold=True)

        tr = table.rows[r]._tr
        trPr = tr.get_or_add_trPr()
        trHeight = trPr.makeelement(qn("w:trHeight"), {qn("w:val"): "700", qn("w:hRule"): "atLeast"})
        trPr.append(trHeight)


def build_doc(participant_id: int, sequence: list[int]) -> Document:
    doc = Document()
    setup_document(doc)
    add_participant_content(doc, participant_id, sequence)
    return doc


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    combined = Document()
    setup_document(combined)

    with SEQUENCES_CSV.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["participant_id"] + [f"n{i+1}" for i in range(SEQUENCE_LENGTH)])

        for participant_id in range(FIRST_ID, LAST_ID + 1):
            sequence = sequence_for(participant_id)

            doc = build_doc(participant_id, sequence)
            doc.save(OUTPUT_DIR / f"participant_{participant_id:03d}.docx")

            add_participant_content(combined, participant_id, sequence)
            if participant_id != LAST_ID:
                combined.add_page_break()

            writer.writerow([participant_id] + sequence)

    combined_path = OUTPUT_DIR / "all_participants.docx"
    combined.save(combined_path)

    count = LAST_ID - FIRST_ID + 1
    print(f"Generated {count} documents in {OUTPUT_DIR}")
    print(f"Wrote combined {count}-page document: {combined_path}")
    print(f"Wrote {SEQUENCES_CSV}")


if __name__ == "__main__":
    main()
